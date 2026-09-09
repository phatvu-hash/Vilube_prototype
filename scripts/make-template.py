"""Dựng file template .xlsx cho team nhập đơn nhập / đơn xuất.

Nguồn dữ liệu là /tmp/td.json do scripts/template-data.ts xuất ra, nên nội dung
mẫu và danh mục trong template luôn khớp với code.
"""
import csv
import io
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

NAVY = "1B3A5B"
HEAD_BG = "E7EDF5"
INPUT_BG = "FFFDE7"   # ô team điền
AUTO_BG = "F1F3F6"    # ô tuỳ chọn, bỏ trống được
FONT = "Arial"

thin = Side(style="thin", color="D6DEE8")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

data = json.load(open(sys.argv[1], encoding="utf-8"))
out_path = sys.argv[2]

# Cột nào bắt buộc, cột nào bỏ trống được — dùng để tô nền và viết chú thích
OPTIONAL = {"GHI_CHU", "SO_LO", "NSX", "HSD", "SO_KIEN", "SL_MOI_KIEN",
            "MA_KIEN_DAU", "DA_NHAN", "SO_LO_NCC", "DVT", "TON_TAI_VI_TRI"}

NOTES = {
    "KHO": "BB = Kho Bao Bì · NVL = Kho Nguyên vật liệu",
    "MA_DON": "Các dòng cùng mã đơn gộp thành một đơn nhập",
    "NGAY_GIAO": "Dạng dd/MM/yyyy",
    "MA_HANG": "Phải có trong tab DANH_MUC và đúng kho",
    "SO_LUONG": "Theo đơn vị cơ sở: CÁI với kho BB, KG với kho NVL",
    "NSX": "Dạng dd/MM/yyyy, bỏ trống được",
    "HSD": "Dạng dd/MM/yyyy, bỏ trống được",
    "SO_KIEN": "Số carton/phuy đã dán tem. Bỏ trống cả ba cột kiện = hàng chưa dán tem",
    "SL_MOI_KIEN": "Số lượng trong một kiện",
    "MA_KIEN_DAU": "Mã kiện đầu tiên, app tự tăng dần. Kho BB cần 7 chữ số (0008083), kho NVL là DrumID (DRM100001)",
    "DA_NHAN": "Số lượng coi như đã nhận sẵn, để có việc Cất hàng mà không phải bấm nhận trước",
    "MA_DON_HANG": "Các dòng cùng mã đơn hàng gộp thành một phiếu soạn",
    "VI_TRI": "Phải có trong tab DANH_MUC và thuộc đúng kho",
    "MA_PALLET": "Pallet ID với kho BB, Drum ID với kho NVL",
    "DVT": "Bỏ trống thì lấy đơn vị cơ sở của mã hàng",
    "TON_TAI_VI_TRI": "Bỏ trống thì hệ thống đặt tồn bằng gấp đôi số lượng yêu cầu",
}


def write_table(ws, csv_text, title_note):
    rows = list(csv.reader(io.StringIO(csv_text)))
    header, body = rows[0], rows[1:]

    for c, name in enumerate(header, start=1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.font = Font(name=FONT, bold=True, size=10, color=NAVY)
        cell.fill = PatternFill("solid", fgColor=HEAD_BG)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
        if name in NOTES:
            cell.comment = None  # chú thích đặt ở tab HUONG_DAN cho dễ đọc
        width = max(len(name) + 2, 12)
        for r in body:
            if c <= len(r):
                width = max(width, min(len(r[c - 1]) + 2, 34))
        ws.column_dimensions[get_column_letter(c)].width = width

    for r, row in enumerate(body, start=2):
        for c, name in enumerate(header, start=1):
            v = row[c - 1] if c <= len(row) else ""
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = Font(name=FONT, size=10)
            cell.fill = PatternFill("solid", fgColor=AUTO_BG if name in OPTIONAL else INPUT_BG)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="center")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{len(body) + 1}"

    note = ws.cell(row=len(body) + 3, column=1, value=title_note)
    note.font = Font(name=FONT, size=9, italic=True, color="5A6270")
    return header, len(body) + 1


wb = Workbook()

# ---------------- DON_NHAP ----------------
ws_in = wb.active
ws_in.title = "DON_NHAP"
head_in, last_in = write_table(
    ws_in,
    data["inboundCsv"],
    "Các dòng trên là đơn mẫu — xoá đi và nhập đơn của bạn. Ô nền vàng là bắt buộc, ô nền xám bỏ trống được.",
)

# ---------------- DON_XUAT ----------------
ws_out = wb.create_sheet("DON_XUAT")
head_out, last_out = write_table(
    ws_out,
    data["outboundCsv"],
    "Các dòng trên là phiếu soạn mẫu — xoá đi và nhập phiếu của bạn. Ô nền vàng là bắt buộc, ô nền xám bỏ trống được.",
)

# ---------------- DANH_MUC ----------------
ws_cat = wb.create_sheet("DANH_MUC")
ws_cat["A1"] = "Danh mục hợp lệ — app chỉ nhận các giá trị dưới đây"
ws_cat["A1"].font = Font(name=FONT, bold=True, size=11, color=NAVY)

blocks = [
    ("MÃ HÀNG", ["Mã hàng", "Tên hàng", "Kho", "Quy cách", "SL/kiện", "Kiện/pallet", "Kg/kiện"],
     [[i["code"], i["name"], i["wh"], i["packing"], i["perCarton"], i["perPallet"], i["kg"]] for i in data["items"]]),
    ("VỊ TRÍ", ["Kho", "Khu vực", "Vị trí"],
     [[l["wh"], l["zone"], l["code"]] for l in data["locations"]]),
    ("LOẠI ĐƠN NHẬP", ["Loại đơn nhập"], [[t] for t in data["inboundTypes"]]),
    ("MÃ KHO", ["Mã kho", "Mã hệ thống", "Tên kho"],
     [[w["kind"], w["code"], w["name"]] for w in data["warehouses"]]),
]

row = 3
item_first_row = None
for title, header, body in blocks:
    c = ws_cat.cell(row=row, column=1, value=title)
    c.font = Font(name=FONT, bold=True, size=10, color=NAVY)
    row += 1
    for ci, h in enumerate(header, start=1):
        hc = ws_cat.cell(row=row, column=ci, value=h)
        hc.font = Font(name=FONT, bold=True, size=10)
        hc.fill = PatternFill("solid", fgColor=HEAD_BG)
        hc.border = BORDER
    row += 1
    if title == "MÃ HÀNG":
        item_first_row = row
    for r in body:
        for ci, v in enumerate(r, start=1):
            vc = ws_cat.cell(row=row, column=ci, value=v)
            vc.font = Font(name=FONT, size=10)
            vc.border = BORDER
        row += 1
    if title == "MÃ HÀNG":
        item_last_row = row - 1
    row += 1

for ci, w in enumerate([14, 42, 8, 18, 12, 14, 12], start=1):
    ws_cat.column_dimensions[get_column_letter(ci)].width = w

# Dropdown mã kho cho hai tab nhập liệu
dv_kho = DataValidation(type="list", formula1='"BB,NVL"', allow_blank=False)
ws_in.add_data_validation(dv_kho)
dv_kho.add(f"A2:A{max(last_in, 200)}")
dv_kho2 = DataValidation(type="list", formula1='"BB,NVL"', allow_blank=False)
ws_out.add_data_validation(dv_kho2)
dv_kho2.add(f"A2:A{max(last_out, 200)}")

# ---------------- HUONG_DAN ----------------
ws_help = wb.create_sheet("HUONG_DAN")
ws_help["A1"] = "HƯỚNG DẪN NHẬP ĐƠN — WMS VILUBE"
ws_help["A1"].font = Font(name=FONT, bold=True, size=14, color=NAVY)

intro = [
    "",
    "Tải file này lên Google Drive rồi mở bằng Google Sheets. Chia sẻ ở mức \"Bất kỳ ai có đường liên kết\" → Người xem,",
    "sau đó gửi link cho đội kỹ thuật để cấu hình. App đọc trực tiếp hai tab DON_NHAP và DON_XUAT.",
    "",
    "Quy ước màu: ô nền vàng là cột bắt buộc, ô nền xám là cột bỏ trống được.",
    "Mỗi dòng là MỘT DÒNG HÀNG, không phải một đơn. Các dòng cùng mã đơn tự gộp thành một đơn.",
    "Dòng nào sai thì app bỏ qua dòng đó và báo rõ tab, số dòng, cột sai; các dòng còn lại vẫn tải bình thường.",
    "Sau khi sửa Sheet, bấm nút Tải lại ở màn hình Chọn kho trong app để lấy bản mới.",
    "",
]
r = 2
for line in intro:
    ws_help.cell(row=r, column=1, value=line).font = Font(name=FONT, size=10)
    r += 1

for tab, header in (("DON_NHAP", head_in), ("DON_XUAT", head_out)):
    t = ws_help.cell(row=r, column=1, value=f"Ý nghĩa các cột — tab {tab}")
    t.font = Font(name=FONT, bold=True, size=11, color=NAVY)
    r += 1
    for h in ("Cột", "Bắt buộc", "Giải thích"):
        hc = ws_help.cell(row=r, column=("Cột", "Bắt buộc", "Giải thích").index(h) + 1, value=h)
        hc.font = Font(name=FONT, bold=True, size=10)
        hc.fill = PatternFill("solid", fgColor=HEAD_BG)
        hc.border = BORDER
    r += 1
    for name in header:
        ws_help.cell(row=r, column=1, value=name).font = Font(name=FONT, size=10)
        ws_help.cell(row=r, column=2, value="" if name in OPTIONAL else "✓").font = Font(name=FONT, size=10)
        ws_help.cell(row=r, column=3, value=NOTES.get(name, "")).font = Font(name=FONT, size=10)
        for ci in (1, 2, 3):
            ws_help.cell(row=r, column=ci).border = BORDER
        r += 1
    r += 1

ws_help.column_dimensions["A"].width = 20
ws_help.column_dimensions["B"].width = 11
ws_help.column_dimensions["C"].width = 96
for row_cells in ws_help.iter_rows(min_col=3, max_col=3):
    for cell in row_cells:
        cell.alignment = Alignment(wrap_text=True, vertical="top")

wb.save(out_path)
print("saved", out_path)
